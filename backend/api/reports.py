from fastapi import APIRouter, Depends, HTTPException, Body
from fastapi.responses import HTMLResponse, FileResponse
from sqlalchemy.orm import Session
from core.database import get_db
from domain import models, schemas
from api.auth import get_current_user
from services.llm import generate_report_content
from services.integrations import fetch_cibil_score, fetch_environmental_risk
import markdown
import tempfile
import os
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
import docx

router = APIRouter()

class ReportUpdate(schemas.ReportBase):
    pass

@router.post("/generate", response_model=schemas.Report)
def generate_report(
    context: str = "",
    client_name: str = "",
    insurance_type: str = "",
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    sources = db.query(models.Source).filter(models.Source.owner_id == current_user.id).all()
    if not sources:
        raise HTTPException(status_code=400, detail="No sources available to generate report. Please upload files first.")

    sources_text = "\n\n".join([f"--- Source: {s.filename} ---\n{s.content}" for s in sources])

    # Enrich with mock API data
    cibil = fetch_cibil_score(client_name or "mock client")
    env = fetch_environmental_risk("mock location")
    enriched_data = f"{sources_text}\n\n--- External API Data ---\nCIBIL Score: {cibil}\nEnvironmental Risk: {env}"

    full_context = f"Client Name: {client_name}\nInsurance Type: {insurance_type}\nUser Instructions: {context}"
    report_content = generate_report_content(enriched_data, full_context)

    db_report = models.Report(
        title=f"{client_name} - {insurance_type} Risk Assessment" if client_name else "AI Risk Assessment Report",
        content=report_content,
        owner_id=current_user.id
    )
    db.add(db_report)
    db.commit()
    db.refresh(db_report)

    return db_report

@router.put("/{report_id}", response_model=schemas.Report)
def update_report(
    report_id: int,
    report_update: ReportUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    report = db.query(models.Report).filter(models.Report.id == report_id, models.Report.owner_id == current_user.id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")

    report.title = report_update.title
    report.content = report_update.content
    db.commit()
    db.refresh(report)
    return report

@router.get("/", response_model=list[schemas.Report])
def get_reports(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    # Return reports ordered by created_at desc to show newest first
    return db.query(models.Report).filter(models.Report.owner_id == current_user.id).order_by(models.Report.created_at.desc()).all()

@router.get("/{report_id}/export/html", response_class=HTMLResponse)
def export_html(report_id: int, db: Session = Depends(get_db)):
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    html_content = markdown.markdown(report.content)
    return HTMLResponse(content=html_content)

@router.get("/{report_id}/export/pdf")
def export_pdf(report_id: int, db: Session = Depends(get_db)):
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")

    fd, path = tempfile.mkstemp(suffix=".pdf")
    os.close(fd)

    c = canvas.Canvas(path, pagesize=letter)
    textobject = c.beginText()
    textobject.setTextOrigin(10, 730)
    textobject.setFont("Helvetica", 10)

    lines = report.content.split('\n')
    for line in lines:
        import textwrap
        wrapped_lines = textwrap.wrap(line, width=100)
        for w_line in wrapped_lines:
            textobject.textLine(w_line)
    c.drawText(textobject)
    c.save()

    return FileResponse(path, filename=f"report_{report_id}.pdf")

@router.get("/{report_id}/export/docx")
def export_docx(report_id: int, db: Session = Depends(get_db)):
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")

    doc = docx.Document()
    doc.add_heading('Insurance Risk Assessment Report', 0)

    for line in report.content.split('\n'):
        if line.strip():
            doc.add_paragraph(line)

    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)

    return FileResponse(path, filename=f"report_{report_id}.docx")
